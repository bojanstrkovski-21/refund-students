import io
from datetime import datetime

from flask import Blueprint, current_app, jsonify, request, session
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from auth import admin_required, login_required
from banks import bank_name
from db import SCHOOL_NAME, get_db

reports_bp = Blueprint('reports', __name__, url_prefix='/api/reports')

MONTHS_MK = [
    '', 'јануари', 'февруари', 'март', 'април', 'мај', 'јуни',
    'јули', 'август', 'септември', 'октомври', 'ноември', 'декември',
]
HEADER_FILL = PatternFill('solid', fgColor='286983')
TOTAL_FILL = PatternFill('solid', fgColor='CFE4E8')
PAID_FILL = PatternFill('solid', fgColor='D7E8D4')
MONEY = '#,##0.00'


def clean(value):
    if value is None:
        return ''
    return str(value).strip()


def period(args):
    year = clean(args.get('year'))
    month = clean(args.get('month'))
    if not year.isdigit() or not (2000 <= int(year) <= 2100):
        return None, None, 'bad_year'
    if not month.isdigit() or not (1 <= int(month) <= 12):
        return None, None, 'bad_month'
    return int(year), int(month), None


def month_name(month):
    return MONTHS_MK[month]


def with_bank(row):
    item = dict(row)
    item['bank_name'] = bank_name(item.get('bank_account'))
    return item


def totals(db, year, month=None):
    if month is None:
        row = db.execute(
            '''SELECT COUNT(*) AS bill_count,
                      COUNT(DISTINCT student_id) AS student_count,
                      ROUND(COALESCE(SUM(amount), 0), 2) AS total
               FROM bills WHERE year = ?''',
            (year,)
        ).fetchone()
    else:
        row = db.execute(
            '''SELECT COUNT(*) AS bill_count,
                      COUNT(DISTINCT student_id) AS student_count,
                      ROUND(COALESCE(SUM(amount), 0), 2) AS total
               FROM bills WHERE year = ? AND month = ?''',
            (year, month)
        ).fetchone()
    return dict(row)


def payment_state(db, year, month, current):
    row = db.execute(
        'SELECT * FROM month_payments WHERE year = ? AND month = ?',
        (year, month)
    ).fetchone()
    if not row:
        return None
    state = dict(row)
    state['stale'] = (
        abs(state['amount_at_payment'] - current['total']) > 0.009
        or state['student_count'] != current['student_count']
        or state['bill_count'] != current['bill_count']
    )
    return state


def students_for_period(db, year, month=None):
    where = 'b.year = ?'
    params = [year]
    if month is not None:
        where += ' AND b.month = ?'
        params.append(month)
    rows = db.execute(f'''
        SELECT s.*,
               COUNT(b.id) AS bill_count,
               ROUND(SUM(b.amount), 2) AS total
        FROM bills b
        JOIN students s ON s.id = b.student_id
        WHERE {where}
        GROUP BY s.id
        ORDER BY s.last_name COLLATE NOCASE, s.first_name COLLATE NOCASE, s.id
    ''', params).fetchall()
    return [with_bank(row) for row in rows]


def bills_for(db, year, month=None, student_id=None):
    clauses = ['b.year = ?']
    params = [year]
    if month is not None:
        clauses.append('b.month = ?')
        params.append(month)
    if student_id is not None:
        clauses.append('b.student_id = ?')
        params.append(student_id)
    rows = db.execute(f'''
        SELECT b.*, s.last_name, s.first_name, s.parent_name, s.class_name, s.parent_embg, s.bank_account
        FROM bills b
        JOIN students s ON s.id = b.student_id
        WHERE {' AND '.join(clauses)}
        ORDER BY s.last_name COLLATE NOCASE, s.first_name COLLATE NOCASE, b.bill_date, b.id
    ''', params).fetchall()
    return [with_bank(row) for row in rows]


@reports_bp.route('/month')
@login_required
def month_summary():
    year, month, error = period(request.args)
    if error:
        return jsonify({'error': error, 'code': error}), 400
    db = get_db()
    current = totals(db, year, month)
    return jsonify({
        'school': SCHOOL_NAME,
        'year': year,
        'month': month,
        'month_name': month_name(month),
        'rows': students_for_period(db, year, month),
        'totals': current,
        'payment': payment_state(db, year, month, current),
    })


@reports_bp.route('/month/bills')
@login_required
def month_bills():
    year, month, error = period(request.args)
    if error:
        return jsonify({'error': error, 'code': error}), 400
    db = get_db()
    current = totals(db, year, month)
    groups = []
    for student in students_for_period(db, year, month):
        groups.append({
            **student,
            'bills': [
                {'id': bill['id'], 'bill_date': bill['bill_date'], 'amount': bill['amount'], 'note': bill['note']}
                for bill in bills_for(db, year, month, student['id'])
            ],
        })
    return jsonify({
        'school': SCHOOL_NAME,
        'year': year,
        'month': month,
        'month_name': month_name(month),
        'groups': groups,
        'totals': current,
        'payment': payment_state(db, year, month, current),
    })


@reports_bp.route('/student-month')
@login_required
def student_month():
    year, month, error = period(request.args)
    if error:
        return jsonify({'error': error, 'code': error}), 400
    student_id = clean(request.args.get('student_id'))
    if not student_id.isdigit():
        return jsonify({'error': 'student_required', 'code': 'student_required'}), 400
    db = get_db()
    student = db.execute('SELECT * FROM students WHERE id = ?', (int(student_id),)).fetchone()
    if not student:
        return jsonify({'error': 'not_found', 'code': 'not_found'}), 404
    bills = bills_for(db, year, month, int(student_id))
    current = totals(db, year, month)
    return jsonify({
        'school': SCHOOL_NAME,
        'year': year,
        'month': month,
        'month_name': month_name(month),
        'student': dict(student),
        'bills': bills,
        'bill_count': len(bills),
        'total': round(sum(bill['amount'] for bill in bills), 2),
        'payment': payment_state(db, year, month, current),
    })


@reports_bp.route('/student')
@login_required
def student_report():
    student_id = clean(request.args.get('student_id'))
    if not student_id.isdigit():
        return jsonify({'error': 'student_required', 'code': 'student_required'}), 400
    year = clean(request.args.get('year'))
    year_value = int(year) if year.isdigit() else None
    db = get_db()
    student = db.execute('SELECT * FROM students WHERE id = ?', (int(student_id),)).fetchone()
    if not student:
        return jsonify({'error': 'not_found', 'code': 'not_found'}), 404
    clauses = ['student_id = ?']
    params = [int(student_id)]
    if year_value:
        clauses.append('year = ?')
        params.append(year_value)
    bills = [dict(row) for row in db.execute(
        f'SELECT * FROM bills WHERE {" AND ".join(clauses)} ORDER BY bill_date, id',
        params
    ).fetchall()]
    months = []
    seen = sorted({(bill['year'], bill['month']) for bill in bills})
    for item_year, item_month in seen:
        item_bills = [bill for bill in bills if bill['year'] == item_year and bill['month'] == item_month]
        current = totals(db, item_year, item_month)
        months.append({
            'year': item_year,
            'month': item_month,
            'month_name': month_name(item_month),
            'bills': item_bills,
            'total': round(sum(bill['amount'] for bill in item_bills), 2),
            'payment': payment_state(db, item_year, item_month, current),
        })
    return jsonify({
        'school': SCHOOL_NAME,
        'student': dict(student),
        'year': year_value,
        'months': months,
        'bill_count': len(bills),
        'total': round(sum(bill['amount'] for bill in bills), 2),
    })


@reports_bp.route('/year')
@login_required
def year_report():
    year = clean(request.args.get('year'))
    if not year.isdigit():
        return jsonify({'error': 'bad_year', 'code': 'bad_year'}), 400
    year = int(year)
    db = get_db()
    rows = db.execute('''
        SELECT s.id, s.last_name, s.first_name, s.parent_name, s.class_name, s.parent_embg, s.bank_account,
               b.month, ROUND(SUM(b.amount), 2) AS total
        FROM bills b
        JOIN students s ON s.id = b.student_id
        WHERE b.year = ?
        GROUP BY s.id, b.month
        ORDER BY s.last_name COLLATE NOCASE, s.first_name COLLATE NOCASE, s.id
    ''', (year,)).fetchall()
    students = {}
    for row in rows:
        item = students.setdefault(row['id'], {
            'id': row['id'],
            'last_name': row['last_name'],
            'first_name': row['first_name'],
            'parent_name': row['parent_name'],
            'class_name': row['class_name'],
            'parent_embg': row['parent_embg'],
            'bank_account': row['bank_account'],
            'months': [0] * 12,
            'total': 0,
        })
        item['months'][row['month'] - 1] = row['total']
        item['total'] = round(item['total'] + row['total'], 2)
    month_rows = []
    for month in range(1, 13):
        current = totals(db, year, month)
        month_rows.append({
            'month': month,
            'month_name': month_name(month),
            'total': current['total'],
            'student_count': current['student_count'],
            'bill_count': current['bill_count'],
            'payment': payment_state(db, year, month, current),
        })
    return jsonify({
        'school': SCHOOL_NAME,
        'year': year,
        'students': list(students.values()),
        'months': month_rows,
        'total': round(sum(item['total'] for item in month_rows), 2),
    })


@reports_bp.route('/month-payment', methods=['POST'])
@admin_required
def mark_month():
    data = request.get_json(silent=True) or {}
    year, month, error = period(data)
    if error:
        return jsonify({'error': error, 'code': error}), 400
    db = get_db()
    current = totals(db, year, month)
    paid = bool(data.get('paid', True))
    if not paid:
        db.execute('DELETE FROM month_payments WHERE year = ? AND month = ?', (year, month))
        db.commit()
        return jsonify({'payment': None, 'totals': current})
    if current['bill_count'] == 0:
        return jsonify({'error': 'empty_month', 'code': 'empty_month'}), 400
    now = datetime.now().strftime('%Y-%m-%d %H:%M')
    db.execute('''
        INSERT INTO month_payments
            (year, month, paid_at, paid_by, amount_at_payment, student_count, bill_count)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(year, month) DO UPDATE SET
            paid_at = excluded.paid_at,
            paid_by = excluded.paid_by,
            amount_at_payment = excluded.amount_at_payment,
            student_count = excluded.student_count,
            bill_count = excluded.bill_count
    ''', (year, month, now, session['username'], current['total'],
          current['student_count'], current['bill_count']))
    db.commit()
    return jsonify({'payment': payment_state(db, year, month, current), 'totals': current})


def text_cell(cell, value):
    cell.value = value or ''
    cell.number_format = '@'
    cell.alignment = Alignment(horizontal='left')


def money_cell(cell, value, bold=False):
    cell.value = float(value or 0)
    cell.number_format = MONEY
    if bold:
        cell.font = Font(bold=True)


def style_header(ws, row, columns):
    for index in range(1, columns + 1):
        cell = ws.cell(row, index)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = HEADER_FILL


def workbook_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    return current_app.response_class(
        buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename={filename}'}
    )


def title_block(ws, title, columns):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws.cell(1, 1, SCHOOL_NAME).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    ws.cell(2, 1, title).font = Font(bold=True, size=12)


@reports_bp.route('/export')
@login_required
def export_report():
    kind = clean(request.args.get('type')) or 'month'
    db = get_db()
    wb = Workbook()
    ws = wb.active
    if kind == 'month':
        year, month, error = period(request.args)
        if error:
            return jsonify({'error': error, 'code': error}), 400
        rows = students_for_period(db, year, month)
        current = totals(db, year, month)
        payment = payment_state(db, year, month, current)
        ws.title = 'Месечен преглед'
        headers = ['#', 'Презиме', 'Име', 'Клас', 'Родител', 'ЕМБГ', 'Трансакциска сметка', 'Банка', 'Сметки', 'Износ']
        title_block(ws, f'Месечен преглед — {month_name(month)} {year}', len(headers))
        status = 'неплатен'
        if payment:
            status = 'платен' + (' (износот е сменет)' if payment['stale'] else '')
        ws.cell(3, 1, f'Статус: {status}').font = Font(italic=True)
        for index, header in enumerate(headers, 1):
            ws.cell(5, index, header)
        style_header(ws, 5, len(headers))
        for index, row in enumerate(rows, 1):
            line = 5 + index
            ws.cell(line, 1, index)
            ws.cell(line, 2, row['last_name'])
            ws.cell(line, 3, row['first_name'])
            ws.cell(line, 4, row['class_name'] or '')
            ws.cell(line, 5, row['parent_name'])
            text_cell(ws.cell(line, 6), row['parent_embg'])
            text_cell(ws.cell(line, 7), row['bank_account'])
            ws.cell(line, 8, row.get('bank_name') or '')
            ws.cell(line, 9, row['bill_count'])
            money_cell(ws.cell(line, 10), row['total'])
        total_row = 6 + len(rows)
        ws.cell(total_row, 5, 'ВКУПНО').font = Font(bold=True)
        ws.cell(total_row, 9, current['bill_count']).font = Font(bold=True)
        money_cell(ws.cell(total_row, 10), current['total'], bold=True)
        ws.cell(total_row, 10).fill = TOTAL_FILL
        widths = [6, 18, 16, 12, 24, 18, 20, 42, 12, 14]
        filename = f'povrat-{year}-{month:02d}.xlsx'
    elif kind == 'month-bills':
        year, month, error = period(request.args)
        if error:
            return jsonify({'error': error, 'code': error}), 400
        ws.title = 'Сметки'
        headers = ['Презиме', 'Име', 'Родител', 'Датум', 'Износ', 'Забелешка']
        title_block(ws, f'Сметки за {month_name(month)} {year}', len(headers))
        for index, header in enumerate(headers, 1):
            ws.cell(4, index, header)
        style_header(ws, 4, len(headers))
        line = 5
        for bill in bills_for(db, year, month):
            ws.cell(line, 1, bill['last_name'])
            ws.cell(line, 2, bill['first_name'])
            ws.cell(line, 3, bill['parent_name'])
            ws.cell(line, 4, '.'.join(reversed(bill['bill_date'].split('-'))))
            money_cell(ws.cell(line, 5), bill['amount'])
            ws.cell(line, 6, bill['note'] or '')
            line += 1
        current = totals(db, year, month)
        ws.cell(line, 3, 'ВКУПНО').font = Font(bold=True)
        money_cell(ws.cell(line, 5), current['total'], bold=True)
        widths = [18, 16, 24, 14, 14, 28]
        filename = f'povrat-smetki-{year}-{month:02d}.xlsx'
    elif kind == 'student-month':
        year, month, error = period(request.args)
        if error:
            return jsonify({'error': error, 'code': error}), 400
        student_id = clean(request.args.get('student_id'))
        if not student_id.isdigit():
            return jsonify({'error': 'student_required', 'code': 'student_required'}), 400
        student = db.execute('SELECT * FROM students WHERE id = ?', (int(student_id),)).fetchone()
        if not student:
            return jsonify({'error': 'not_found', 'code': 'not_found'}), 404
        ws.title = 'Ученик'
        headers = ['#', 'Датум', 'Износ', 'Забелешка']
        title_block(ws, f'{student["last_name"]} {student["first_name"]} — {month_name(month)} {year}', len(headers))
        ws.cell(3, 1, f'Родител: {student["parent_name"]}')
        embg = ws.cell(3, 3, student['parent_embg'])
        text_cell(embg, student['parent_embg'])
        account = ws.cell(3, 4, student['bank_account'])
        text_cell(account, student['bank_account'])
        ws.cell(4, 1, bank_name(student['bank_account']))
        for index, header in enumerate(headers, 1):
            ws.cell(5, index, header)
        style_header(ws, 5, len(headers))
        bills = bills_for(db, year, month, int(student_id))
        for index, bill in enumerate(bills, 1):
            line = 5 + index
            ws.cell(line, 1, index)
            ws.cell(line, 2, '.'.join(reversed(bill['bill_date'].split('-'))))
            money_cell(ws.cell(line, 3), bill['amount'])
            ws.cell(line, 4, bill['note'] or '')
        total_row = 6 + len(bills)
        ws.cell(total_row, 2, 'ВКУПНО').font = Font(bold=True)
        money_cell(ws.cell(total_row, 3), sum(bill['amount'] for bill in bills), bold=True)
        ws.cell(total_row, 3).fill = TOTAL_FILL
        widths = [6, 14, 14, 28]
        filename = f'povrat-ucenik-{year}-{month:02d}.xlsx'
    elif kind == 'student':
        student_id = clean(request.args.get('student_id'))
        if not student_id.isdigit():
            return jsonify({'error': 'student_required', 'code': 'student_required'}), 400
        student = db.execute('SELECT * FROM students WHERE id = ?', (int(student_id),)).fetchone()
        if not student:
            return jsonify({'error': 'not_found', 'code': 'not_found'}), 404
        year = clean(request.args.get('year'))
        year_value = int(year) if year.isdigit() else None
        clauses = ['student_id = ?']
        params = [int(student_id)]
        if year_value:
            clauses.append('year = ?')
            params.append(year_value)
        bills = db.execute(
            f'SELECT * FROM bills WHERE {" AND ".join(clauses)} ORDER BY bill_date, id',
            params
        ).fetchall()
        ws.title = 'Ученик'
        headers = ['Датум', 'Месец', 'Износ', 'Забелешка']
        period_label = str(year_value) if year_value else 'сите години'
        title_block(ws, f'{student["last_name"]} {student["first_name"]} — {period_label}', len(headers))
        ws.cell(3, 1, f'Родител: {student["parent_name"]}')
        text_cell(ws.cell(3, 3), student['parent_embg'])
        text_cell(ws.cell(3, 4), student['bank_account'])
        for index, header in enumerate(headers, 1):
            ws.cell(5, index, header)
        style_header(ws, 5, len(headers))
        for index, bill in enumerate(bills, 1):
            line = 5 + index
            ws.cell(line, 1, '.'.join(reversed(bill['bill_date'].split('-'))))
            ws.cell(line, 2, f'{month_name(bill["month"])} {bill["year"]}')
            money_cell(ws.cell(line, 3), bill['amount'])
            ws.cell(line, 4, bill['note'] or '')
        total_row = 6 + len(bills)
        ws.cell(total_row, 2, 'ВКУПНО').font = Font(bold=True)
        money_cell(ws.cell(total_row, 3), sum(bill['amount'] for bill in bills), bold=True)
        widths = [14, 20, 14, 28]
        filename = 'povrat-ucenik.xlsx'
    elif kind == 'year':
        year = clean(request.args.get('year'))
        if not year.isdigit():
            return jsonify({'error': 'bad_year', 'code': 'bad_year'}), 400
        year = int(year)
        payload_rows = db.execute('''
            SELECT s.last_name, s.first_name, s.class_name, s.parent_name, b.month, ROUND(SUM(b.amount), 2) AS total
            FROM bills b JOIN students s ON s.id = b.student_id
            WHERE b.year = ?
            GROUP BY s.id, b.month
            ORDER BY s.last_name COLLATE NOCASE, s.first_name COLLATE NOCASE, s.id
        ''', (year,)).fetchall()
        students = {}
        for row in payload_rows:
            key = (row['last_name'], row['first_name'], row['class_name'] or '', row['parent_name'])
            item = students.setdefault(key, [0] * 12)
            item[row['month'] - 1] = row['total']
        ws.title = 'Година'
        headers = ['Презиме', 'Име', 'Клас', 'Родител'] + [month_name(i)[:3] for i in range(1, 13)] + ['Вкупно']
        title_block(ws, f'Годишен извештај {year}', len(headers))
        for index, header in enumerate(headers, 1):
            ws.cell(4, index, header)
        style_header(ws, 4, len(headers))
        for month_index in range(1, 13):
            current = totals(db, year, month_index)
            if payment_state(db, year, month_index, current):
                cell = ws.cell(4, 4 + month_index)
                cell.fill = PAID_FILL
                cell.font = Font(bold=True)
        line = 5
        month_totals = [0] * 12
        for (last_name, first_name, class_name, parent_name), amounts in students.items():
            ws.cell(line, 1, last_name)
            ws.cell(line, 2, first_name)
            ws.cell(line, 3, class_name)
            ws.cell(line, 4, parent_name)
            for index, amount in enumerate(amounts):
                money_cell(ws.cell(line, 5 + index), amount)
                month_totals[index] += amount
            money_cell(ws.cell(line, 17), sum(amounts), bold=True)
            line += 1
        ws.cell(line, 4, 'ВКУПНО').font = Font(bold=True)
        for index, amount in enumerate(month_totals):
            money_cell(ws.cell(line, 5 + index), amount, bold=True)
        money_cell(ws.cell(line, 17), sum(month_totals), bold=True)
        widths = [16, 14, 10, 20] + [11] * 12 + [13]
        filename = f'povrat-{year}.xlsx'
    else:
        return jsonify({'error': 'bad_report', 'code': 'bad_report'}), 400
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = SCHOOL_NAME
    ws.print_title_rows = '1:4'
    ws.page_setup.paperSize = 9
    return workbook_response(wb, filename)
